"""Build the standalone HTML, Excel workbook, and CSV from the reviewed Markdown.

Run: python literature_review/build_review_artifacts.py
Dependencies: markdown, beautifulsoup4, openpyxl
"""
from pathlib import Path
import csv
import math
import re
import unicodedata

import markdown
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '2023plus_calibration_review.md'
raw = SOURCE.read_text(encoding='utf-8')
references = dict(re.findall(r'^\[\^([^\]]+)\]: (.+)$', raw, re.M))


def plain(text):
    return BeautifulSoup(markdown.markdown(text), 'html.parser').get_text(' ', strip=True)


urls = {key: re.search(r'\]\((https?://[^)]+)\)', value).group(1)
        for key, value in references.items()}
rendered = markdown.markdown(raw, extensions=['tables', 'footnotes', 'toc'],
                             extension_configs={'toc': {'toc_depth': '2-3'}})
soup = BeautifulSoup(rendered, 'html.parser')
for table in list(soup.find_all('table')):
    wrapper = soup.new_tag('div', attrs={'class': 'table-scroll'})
    table.wrap(wrapper)
css = '''
:root { color-scheme:light; font-family:Arial,"Noto Sans KR",sans-serif; color:#20252b; }
body { margin:0; background:white; line-height:1.75; }
main { max-width:1200px; margin:auto; padding:38px 30px 80px; }
h1 { font-size:2rem; line-height:1.35; margin:0 0 28px; }
h2 { margin-top:56px; padding-top:14px; border-top:2px solid #343a40; font-size:1.45rem; }
h3 { margin-top:36px; font-size:1.12rem; }
a { color:#18588a; text-underline-offset:3px; }
p { max-width:1050px; }
.table-scroll { overflow-x:auto; margin:20px 0; }
table { width:100%; border-collapse:collapse; font-size:.87rem; line-height:1.6; }
th { background:#30363d; color:white; font-weight:600; text-align:left; }
td,th { padding:10px 12px; border-bottom:1px solid #d7dce1; vertical-align:top; min-width:85px; }
tr:nth-child(even) td { background:#f5f6f7; }
td:first-child { min-width:125px; }
code { background:#f1f2f3; padding:1px 4px; }
.footnote { font-size:.86rem; overflow-wrap:anywhere; }
.toc { columns:2; border-bottom:1px solid #ddd; padding-bottom:20px; }
.toc ul { padding-left:20px; }
.toc li { margin:4px 0; }
nav { margin-bottom:24px; font-size:.9rem; }
@media(max-width:700px) { main { padding:22px 16px; } .toc { columns:1; } h1 { font-size:1.6rem; } }
@media print { main { max-width:none; padding:0; } nav,.toc { display:none; }
 h2,h3 { break-after:avoid; } tr { break-inside:avoid; } .table-scroll { overflow:visible; }
 th { background:#eee; color:black; } table { font-size:8pt; } td,th { min-width:0; padding:5px; } }
'''
toc_renderer = markdown.Markdown(extensions=['toc', 'tables', 'footnotes'],
                                 extension_configs={'toc': {'toc_depth': '2-3'}})
toc_renderer.convert(raw)
title = soup.h1.extract()
html = ('<!doctype html><html lang="ko"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<title>2023년 이후 캘리브레이션 실측 문헌 비교</title>'
        f'<style>{css}</style></head><body><main>{title}'
        '<nav><a href="calibration_papers_real_measurements.xlsx">비교 엑셀</a> · '
        '<a href="paper_comparison.csv">논문 목록 CSV</a> · '
        '<a href="2023plus_calibration_review.md">Markdown 원본</a></nav>'
        f'{toc_renderer.toc}{soup}</main></body></html>')
(ROOT / '2023plus_calibration_review.html').write_text(html, encoding='utf-8')

wb = Workbook()
wb.remove(wb.active)
fill = PatternFill('solid', fgColor='30363D')
alternate = PatternFill('solid', fgColor='F1F3F5')


def width(text):
    return sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in str(text))


def sheet(name, headers, rows, source_note='', context=''):
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.showGridLines = False
    ws.append(headers)
    for values in rows:
        ws.append(values)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 42
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='맑은 고딕', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = alternate
            if source_note:
                cell.comment = Comment(source_note + '\n\n' + context, 'Sources')
            if isinstance(cell.value, str) and cell.value.startswith('https://'):
                cell.hyperlink = cell.value
                cell.font = Font(name='맑은 고딕', size=10, color='18588A', underline='single')
    for col in ws.columns:
        col_width = min(65, max(15, max(width(c.value or '') for c in col) + 2))
        ws.column_dimensions[col[0].column_letter].width = col_width
    for row in ws.iter_rows(min_row=2):
        lines = max(math.ceil(width(c.value or '') /
                    max(10, ws.column_dimensions[c.column_letter].width - 2)) for c in row)
        ws.row_dimensions[row[0].row].height = min(350, max(34, 16 * lines + 10))
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = '1:1'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


sheet('읽는법', ['항목', '설명'], [
    ['범위', '2023-01-01~2026-09-09. 논문 25편 검토. 실제 장비·실촬영 결과만 수록.'],
    ['제외', '시뮬레이션, 인위적 시간 offset 실험, 실측 조건 미확인 초록 수치는 정량 비교 제외.'],
    ['S1', '해당 논문 안의 동일 실측 조건·지표에서 확인되는 우위. 전 분야 현재 최고가 아님.'],
    ['S2 / B / E', '최신 후보 / baseline·선행연구 / 평가·도구.'],
    ['GT', '실측 여부와 GT 독립성은 별개. 재투영·폐루프·반복성·targeting·pose GT 오차를 분리.'],
    ['미확인 / 미보고', '원문 접근·검증 부족 / 확인한 논문에 해당 결과 없음. 0으로 해석하지 말 것.'],
    ['관련성', '5/5는 이 프로젝트의 직접 비교·불확실성·큐브 형상·독립 평가에 중요.'],
    ['단위', '원문 px, mm, cm, rad, deg, mm²를 유지. 환산은 별도 표기.'],
    ['출처', '각 데이터 셀 메모에 원문 및 표 주변 해석 기록. Sources 시트에 전체 참고문헌.'],
    ['사용법', '논문목록 필터로 분야·SOTA·관련성을 선택. 상세 표는 논문 ID별 시트.'],
    ['해석', '2023plus_calibration_review.html에 실측 조건, 비교 한계, 프로젝트 분석 전체 수록.'],
])

table_counts = {}
index_rows = None
section_names = {'1.0': '신규물체_2~3mm비교', '1.1': '프로젝트구성', '1.2': '문서차이', '2.': 'SOTA표시',
                 '7.': 'SOTA판단', '8.': '지표정의', '8.1': '우리실측',
                 '8.2': '우선비교군', '9.': '추가후보'}
for table in soup.find_all('table'):
    heading = table.find_previous(['h2', 'h3']).get_text(' ', strip=True)
    rows = [[c.get_text(' ', strip=True) for c in tr.find_all(['th', 'td'])]
            for tr in table.find_all('tr')]
    headers, values = rows[0], rows[1:]
    paper_match = re.match(r'([CRD]\d+)\.', heading)
    key = paper_match.group(1) if paper_match else None
    context_parts = []
    if key:
        h = table.find_previous(['h2', 'h3'])
        for sibling in h.next_siblings:
            if getattr(sibling, 'name', None) in ('h2', 'h3'):
                break
            if getattr(sibling, 'name', None) == 'p':
                context_parts.append(sibling.get_text(' ', strip=True))
    if headers[0] == 'ID':
        name = '논문목록'
        headers = headers + ['원문 링크']
        values = [row + [urls[row[0]]] for row in values]
        index_rows = (headers, values)
    elif key:
        name = key + '_실측결과'
    elif headers[0] == '실측 근거':
        name = '실측근거구분'
    else:
        name = next((v for k, v in section_names.items() if heading.startswith(k)), '범위')
    count = table_counts.get(name, 0) + 1
    table_counts[name] = count
    if count > 1:
        name += '_' + str(count)
    ws = sheet(name, headers, values,
               plain(references[key]) if key else '', '\n\n'.join(context_parts))
    if name == '논문목록':
        for row in ws.iter_rows(min_row=2):
            note = plain(references[row[0].value])
            for cell in row:
                cell.comment = Comment(note, 'Sources')

for key, status in [('C5', '실측 조건별 수치 미확인. 초록 수치는 결과 비교 제외.'),
                    ('C7', 'JOSS 논문에서 정량 비교 성능 미보고.'),
                    ('R4', '이론·코드 확인. 실측 실험표 원문 미확인.'),
                    ('R7', '실제 xArm targeting 약 4 mm. Baxter PCK3D @2cm 0.15 / @5cm 0.80 (3 views).')]:
    sheet(key + '_실측메모', ['논문', '실측 결과 / 상태', '원문'],
          [[key, status, urls[key]]], plain(references[key]))

sheet('Sources', ['ID', '서지·원문 위치·확인 범위', '원문 URL'],
      [[key, plain(value), urls[key]] for key, value in references.items()])
wb._sheets.sort(key=lambda ws: (0 if ws.title == '읽는법' else
                              1 if ws.title == '논문목록' else
                              2 if ws.title == '실측근거구분' else
                              4 if ws.title == 'Sources' else 3))
target = ROOT / 'calibration_papers_real_measurements.xlsx'
wb.save(target)
assert index_rows is not None and len(index_rows[1]) == 25
with (ROOT / 'paper_comparison.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(index_rows[0])
    writer.writerows(index_rows[1])
loaded = load_workbook(target)
assert loaded['논문목록'].max_row == 26
assert len(set(row[0] for row in index_rows[1])) == 25
assert not re.search(r'\|\s*합성|\|\s*Synthetic|검증 메모|easyhec2', raw)
defined = set(references)
used = set(re.findall(r'\[\^([^\]]+)\](?!:)', raw))
assert used <= defined
for link in soup.select('a[href]'):
    href = link['href']
    if href.startswith('../'):
        assert (ROOT / href).exists(), href
print(f'Validated: 25 papers, {len(loaded.sheetnames)} sheets, {len(references)} sources.')
print('Created HTML, XLSX, CSV from the reviewed Markdown.')
